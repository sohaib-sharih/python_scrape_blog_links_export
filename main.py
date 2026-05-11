import sys
import time
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from selenium.webdriver.common.by import By
from seleniumbase import Driver

TARGET_URL = "https://www.smythson.com/uk/smythson-stories.html"
OUTPUT_FILE = "smythson_blog_posts.xlsx"
DEBUG_FILE = "debug.html"

CLOUDFLARE_KEYWORDS = [
    "security verification",
    "verifying you are not a bot",
    "cloudflare",
    "checking your browser",
    "just a moment",
    "ray id:",
]


def main():
    driver = Driver(uc=True, headless=False)
    driver.set_window_size(1280, 800)

    try:
        print(f"Navigating to {TARGET_URL} ...")
        driver.get(TARGET_URL)

        print("\nA browser window has opened.")
        print("If a Cloudflare / human verification page appears, please:")
        print("  1. Wait for the checkbox to load")
        print("  2. Click it to verify you are human")
        print("  3. Wait for the actual blog page to fully load")
        input("Press Enter AFTER the blog listing page is visible ...")

        while is_cloudflare_blocked(driver):
            print("\nCloudflare verification page is still showing.")
            print("Please complete the verification in the browser and wait for the real blog page to appear.")
            input("Press Enter again after the blog page is visible ...")

        time.sleep(3)

        print("Scrolling page to load all content ...")
        scroll_page(driver)

        print("Scraping blog posts ...")
        posts = extract_blog_posts(driver)

        if not posts:
            print("No blog posts found. Saving page HTML to debug.html for inspection ...")
            with open(DEBUG_FILE, "w", encoding="utf-8") as f:
                f.write(driver.page_source)
            print(f"Saved {DEBUG_FILE} — I can use this to fix the selectors.")
        else:
            save_to_excel(posts)
            print(f"\nDone! Found {len(posts)} blog posts.")
            print(f"Results saved to: {OUTPUT_FILE}")

    finally:
        driver.quit()


def is_cloudflare_blocked(driver):
    body = driver.find_element(By.TAG_NAME, "body").text.lower()
    return any(kw in body for kw in CLOUDFLARE_KEYWORDS)


def scroll_page(driver):
    last_height = driver.execute_script("return document.body.scrollHeight")
    for _ in range(6):
        driver.execute_script(f"window.scrollTo(0, {last_height});")
        time.sleep(1.5)
        new_height = driver.execute_script("return document.body.scrollHeight")
        if new_height == last_height:
            break
        last_height = new_height
    driver.execute_script("window.scrollTo(0, 0);")
    time.sleep(1)


def extract_blog_posts(driver):
    posts = []
    seen = set()

    main_content = find_main_content(driver)

    if main_content:
        stories = main_content.find_elements(By.CSS_SELECTOR, ".jour-land")
        for story in stories:
            title_el = story.find_elements(By.CSS_SELECTOR, ".jorn-blog-name")
            if title_el:
                href = title_el[0].get_attribute("href")
                title = title_el[0].text.strip()
                if href and title:
                    key = (href, title)
                    if key not in seen:
                        seen.add(key)
                        posts.append({"title": title, "url": href})

    if not posts:
        container = main_content or driver.find_element(By.TAG_NAME, "body")
        all_links = container.find_elements(By.TAG_NAME, "a")
        for el in all_links:
            href = el.get_attribute("href")
            title = el.text.strip()
            if (
                href
                and title
                and href not in ("#", "", "/")
                and not href.startswith("javascript:")
                and len(title) > 5
            ):
                full_url = urljoin(TARGET_URL, href)
                key = (full_url, title)
                if key not in seen:
                    seen.add(key)
                    posts.append({"title": title, "url": full_url})

    return posts


def find_main_content(driver):
    for selector in ["main", "[role='main']", ".l-body-page_main"]:
        elements = driver.find_elements(By.CSS_SELECTOR, selector)
        if elements:
            return elements[0]
    return None


def save_to_excel(posts):
    wb = Workbook()
    ws = wb.active
    ws.title = "Blog Posts"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

    ws["A1"] = "Blog Title"
    ws["B1"] = "Blog Link"
    for cell in (ws["A1"], ws["B1"]):
        cell.font = header_font
        cell.fill = header_fill

    for i, post in enumerate(posts, start=2):
        ws[f"A{i}"] = post["title"]
        ws[f"B{i}"] = post["url"]

    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 80

    wb.save(OUTPUT_FILE)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nAborted by user.")
        sys.exit(1)
